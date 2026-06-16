"""Unit tests for qa_signoff.docx_builder — build_document."""

from pathlib import Path

import docx
import pytest

from aiflow_demo.qa_signoff.docx_builder import CONFIGURED_GATES, build_document
from aiflow_demo.qa_signoff.models import (
    FailureDetail,
    QaDocument,
    RegressionResult,
    SignoffMeta,
    TestCase,
    TestSection,
)


def _make_meta(tag: str = "v1.0.0", sha: str = "abc1234") -> SignoffMeta:
    return SignoffMeta(
        project="AIFlowTesting",
        branch="feature/qa-signoff-doc",
        tag=tag,
        sha=sha,
        generated_at="2026-06-16T12:00:00+00:00",
    )


def _make_cases() -> tuple[TestCase, ...]:
    return (
        TestCase(1, "Cart adds item", "unit", "tests/test_cart.py", "new cart returned"),
        TestCase(2, "Cart rejects blank name", "unit", "tests/test_cart.py", "CartError"),
    )


def _make_section() -> TestSection:
    return TestSection(
        datetime="2026-06-10 09:00:00",
        title="First task",
        intro="Section intro text.",
        cases=_make_cases(),
    )


def _make_regression_pass() -> RegressionResult:
    return RegressionResult(
        total=3,
        passed=3,
        failed=0,
        skipped=0,
        errors=0,
        duration_s=0.42,
        coverage_pct=87.5,
        failures=(),
        available=True,
    )


def _make_regression_with_failures() -> RegressionResult:
    return RegressionResult(
        total=3,
        passed=2,
        failed=1,
        skipped=0,
        errors=0,
        duration_s=0.5,
        coverage_pct=85.0,
        failures=(FailureDetail("tests.test_x.test_fail", "assert 1 == 2"),),
        available=True,
    )


def _make_doc(
    regression: RegressionResult | None = None,
    sections: tuple[TestSection, ...] | None = None,
) -> QaDocument:
    return QaDocument(
        meta=_make_meta(),
        sections=sections if sections is not None else (_make_section(),),
        regression=regression if regression is not None else _make_regression_pass(),
        gates=CONFIGURED_GATES,
    )


class TestFileSaved:
    def test_file_saved_and_reopenable(self, tmp_path: Path):
        path = build_document(_make_doc(), tmp_path)
        assert path.exists()
        reopened = docx.Document(str(path))
        assert reopened is not None

    def test_creates_missing_output_directory(self, tmp_path: Path):
        nested = tmp_path / "does" / "not" / "exist"
        path = build_document(_make_doc(), nested)
        assert path.exists()
        assert path.parent == nested

    def test_filename_contains_tag(self, tmp_path: Path):
        path = build_document(_make_doc(), tmp_path)
        assert "v1.0.0" in path.name

    def test_filename_sanitizes_unsafe_chars(self, tmp_path: Path):
        meta = _make_meta(tag="v1.0.0/rc:1")
        doc = QaDocument(
            meta=meta,
            sections=(_make_section(),),
            regression=_make_regression_pass(),
            gates=CONFIGURED_GATES,
        )
        path = build_document(doc, tmp_path)
        assert "/" not in path.name
        assert ":" not in path.name


class TestTitleAndMeta:
    def _all_text(self, path: Path) -> str:
        d = docx.Document(str(path))
        return "\n".join(p.text for p in d.paragraphs)

    def test_title_present(self, tmp_path: Path):
        path = build_document(_make_doc(), tmp_path)
        assert "QA Sign-off Report" in self._all_text(path)

    def test_project_present(self, tmp_path: Path):
        path = build_document(_make_doc(), tmp_path)
        assert "AIFlowTesting" in self._all_text(path)

    def test_branch_present(self, tmp_path: Path):
        path = build_document(_make_doc(), tmp_path)
        assert "feature/qa-signoff-doc" in self._all_text(path)

    def test_tag_present(self, tmp_path: Path):
        path = build_document(_make_doc(), tmp_path)
        assert "v1.0.0" in self._all_text(path)

    def test_sha_present(self, tmp_path: Path):
        path = build_document(_make_doc(), tmp_path)
        assert "abc1234" in self._all_text(path)


class TestResultSummary:
    def _all_text(self, path: Path) -> str:
        d = docx.Document(str(path))
        return "\n".join(p.text for p in d.paragraphs)

    def test_verdict_present(self, tmp_path: Path):
        path = build_document(_make_doc(), tmp_path)
        assert "PASS" in self._all_text(path)

    def test_coverage_metric_present(self, tmp_path: Path):
        path = build_document(_make_doc(), tmp_path)
        assert "87.5" in self._all_text(path)

    def test_all_gate_labels_present(self, tmp_path: Path):
        path = build_document(_make_doc(), tmp_path)
        text = self._all_text(path)
        for gate in CONFIGURED_GATES:
            assert gate in text, f"Gate label missing: {gate}"


class TestCatalogTable:
    def test_one_table_per_section(self, tmp_path: Path):
        sections = (_make_section(), _make_section())
        doc = _make_doc(sections=sections)
        path = build_document(doc, tmp_path)
        d = docx.Document(str(path))
        catalog_tables = [
            t for t in d.tables
            if t.rows and any(
                cell.text.strip() in {"Test case", "#"} for cell in t.rows[0].cells
            )
        ]
        assert len(catalog_tables) == 2

    def test_table_header_row_has_correct_columns(self, tmp_path: Path):
        path = build_document(_make_doc(), tmp_path)
        d = docx.Document(str(path))
        catalog_table = next(
            t for t in d.tables
            if t.rows and any(
                cell.text.strip() == "Test case" for cell in t.rows[0].cells
            )
        )
        headers = [c.text.strip() for c in catalog_table.rows[0].cells]
        assert "#" in headers
        assert "Test case" in headers
        assert "Type" in headers
        assert "Test file" in headers
        assert "Expected result" in headers

    def test_table_row_count_equals_cases_plus_one(self, tmp_path: Path):
        path = build_document(_make_doc(), tmp_path)
        d = docx.Document(str(path))
        catalog_table = next(
            t for t in d.tables
            if t.rows and any(
                cell.text.strip() == "Test case" for cell in t.rows[0].cells
            )
        )
        assert len(catalog_table.rows) == len(_make_cases()) + 1


class TestRegressionSection:
    def test_regression_table_present_when_available(self, tmp_path: Path):
        path = build_document(_make_doc(regression=_make_regression_pass()), tmp_path)
        d = docx.Document(str(path))
        all_text = "\n".join(p.text for p in d.paragraphs)
        assert "3" in all_text

    def test_unavailable_note_when_not_available(self, tmp_path: Path):
        doc = _make_doc(regression=RegressionResult.unavailable())
        path = build_document(doc, tmp_path)
        d = docx.Document(str(path))
        all_text = "\n".join(p.text for p in d.paragraphs)
        assert "unavailable" in all_text.lower()

    def test_failures_listed_when_present(self, tmp_path: Path):
        doc = _make_doc(regression=_make_regression_with_failures())
        path = build_document(doc, tmp_path)
        d = docx.Document(str(path))
        all_text = "\n".join(p.text for p in d.paragraphs)
        assert "tests.test_x.test_fail" in all_text
        assert "assert 1 == 2" in all_text


class TestSignoffBlock:
    def test_signoff_table_has_three_roles(self, tmp_path: Path):
        path = build_document(_make_doc(), tmp_path)
        d = docx.Document(str(path))
        signoff_table = next(
            (
                t for t in d.tables
                if any(
                    "QA Lead" in c.text for row in t.rows for c in row.cells
                )
            ),
            None,
        )
        assert signoff_table is not None
        roles = [row.cells[0].text.strip() for row in signoff_table.rows]
        assert "QA Lead" in roles
        assert "Dev Lead" in roles
        assert "Product Owner" in roles

    def test_signoff_table_has_four_columns(self, tmp_path: Path):
        path = build_document(_make_doc(), tmp_path)
        d = docx.Document(str(path))
        signoff_table = next(
            t for t in d.tables
            if any("QA Lead" in c.text for row in t.rows for c in row.cells)
        )
        assert len(signoff_table.columns) == 4

    def test_release_decision_paragraph_present(self, tmp_path: Path):
        path = build_document(_make_doc(), tmp_path)
        d = docx.Document(str(path))
        all_text = "\n".join(p.text for p in d.paragraphs)
        assert "Release decision" in all_text

    def test_signoff_table_has_header_row(self, tmp_path: Path):
        path = build_document(_make_doc(), tmp_path)
        d = docx.Document(str(path))
        signoff_table = next(
            t for t in d.tables
            if any("QA Lead" in c.text for row in t.rows for c in row.cells)
        )
        header = [c.text.strip() for c in signoff_table.rows[0].cells]
        assert header == ["Role", "Name", "Signature", "Date"]


class TestDeterminism:
    def test_build_twice_identical_extracted_text(self, tmp_path: Path):
        doc = _make_doc()
        path1 = tmp_path / "run1"
        path1.mkdir()
        path2 = tmp_path / "run2"
        path2.mkdir()
        p1 = build_document(doc, path1)
        p2 = build_document(doc, path2)
        d1 = docx.Document(str(p1))
        d2 = docx.Document(str(p2))
        text1 = "\n".join(p.text for p in d1.paragraphs)
        text2 = "\n".join(p.text for p in d2.paragraphs)
        assert text1 == text2
