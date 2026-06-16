"""Tests for qa_testcases.render — render_document layout."""

from pathlib import Path

import docx

from aiflow_demo.qa_testcases.models import ModuleCoverage, TestCaseRow, TestPlan
from aiflow_demo.qa_testcases.render import (
    CASE_COLUMNS,
    COVERAGE_COLUMNS,
    RenderMeta,
    render_document,
)


def _plan() -> TestPlan:
    return TestPlan(
        project="DemoApp",
        summary="Plan for the demo.",
        assumptions=("Codes are WELCOME10/VIP20.", "Cart is immutable."),
        coverage=(
            ModuleCoverage("Cart", 2, 4, 2, 1, 1, 0),
            ModuleCoverage("Orders", 1, 2, 1, 0, 0, 1),
        ),
        test_cases=(
            TestCaseRow("US-1", "Adds", "Add valid", "1. add", "x=1", "Positive"),
            TestCaseRow("US-1", "Rejects", "Blank", "1. add ''", "x=''", "Negative"),
        ),
    )


def _meta() -> RenderMeta:
    return RenderMeta(
        branch="v1.0.0", generated_at="2026-06-16 12:00 UTC", date_stamp="20260616"
    )


def _texts(path: Path) -> str:
    d = docx.Document(str(path))
    return "\n".join(p.text for p in d.paragraphs)


def test_creates_file_and_missing_dir(tmp_path: Path):
    out = tmp_path / "nested" / "qa"
    path = render_document(_plan(), _meta(), out)
    assert path.exists()
    assert path.parent == out


def test_filename_follows_convention(tmp_path: Path):
    path = render_document(_plan(), _meta(), tmp_path)
    assert path.name == "Test_Cases_DemoApp_20260616.docx"


def test_title_and_header_present(tmp_path: Path):
    text = _texts(render_document(_plan(), _meta(), tmp_path))
    assert "Test Case Document — DemoApp" in text
    assert "Branch: v1.0.0" in text


def test_coverage_table_has_total_row(tmp_path: Path):
    path = render_document(_plan(), _meta(), tmp_path)
    d = docx.Document(str(path))
    cov = next(t for t in d.tables if t.rows[0].cells[0].text == "Module")
    assert [c.text for c in cov.rows[0].cells] == list(COVERAGE_COLUMNS)
    total = cov.rows[-1]
    assert total.cells[0].text == "TOTAL"
    assert total.cells[2].text == "6"  # 4 + 2 total TCs


def test_case_table_columns_and_blank_result(tmp_path: Path):
    path = render_document(_plan(), _meta(), tmp_path)
    d = docx.Document(str(path))
    case = next(t for t in d.tables if t.rows[0].cells[0].text == "User Story")
    assert [c.text for c in case.rows[0].cells] == list(CASE_COLUMNS)
    assert len(case.rows) == 3  # header + 2 cases
    assert case.rows[1].cells[-1].text == ""  # Test Result left blank for QA


def test_signoff_fields_and_signatures(tmp_path: Path):
    path = render_document(_plan(), _meta(), tmp_path)
    d = docx.Document(str(path))
    text = _texts(path)
    assert "QA Sign-Off" in text
    assert "QA Engineer Signature:" in text
    assert "QA Lead / Manager Approval:" in text
    fields = next(t for t in d.tables if t.rows[0].cells[0].text == "QA Engineer Name")
    labels = [r.cells[0].text for r in fields.rows]
    assert "Passed" in labels
    assert "Failed" in labels
