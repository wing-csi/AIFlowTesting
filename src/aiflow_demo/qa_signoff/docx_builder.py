"""Build a QA sign-off Word document from a QaDocument model."""

from __future__ import annotations

import re
from pathlib import Path

import docx

from .models import QaDocument, RegressionResult, SignoffMeta, TestSection

CONFIGURED_GATES: tuple[str, ...] = (
    "tests",
    "coverage >=80%",
    "ruff lint",
    "CodeQL",
    "security-scan",
)

CATALOG_COLUMNS: tuple[str, ...] = (
    "#",
    "Test case",
    "Type",
    "Test file",
    "Expected result",
)

SIGNOFF_ROLES: tuple[str, ...] = ("QA Lead", "Dev Lead", "Product Owner")

SIGNOFF_COLUMNS: tuple[str, ...] = ("Role", "Name", "Signature", "Date")

GRID_STYLE = "Table Grid"

_UNSAFE_CHARS_RE = re.compile(r'[<>:"/\\|?*]')


def _output_filename(meta: SignoffMeta) -> str:
    """Produce a filesystem-safe filename for the sign-off document."""
    safe_tag = _UNSAFE_CHARS_RE.sub("-", meta.tag)
    return f"qa-signoff-{safe_tag}.docx"


def _fill_table_row(row: docx.table.Row, values: tuple[str, ...]) -> None:
    """Write string values into each cell of a table row."""
    for cell, value in zip(row.cells, values, strict=False):
        cell.text = value


def _add_title_header(document: docx.Document, meta: SignoffMeta) -> None:
    """Add the report title and metadata block."""
    document.add_heading("QA Sign-off Report", level=0)
    document.add_paragraph(f"Project: {meta.project}")
    document.add_paragraph(f"Branch: {meta.branch}")
    document.add_paragraph(f"Tag: {meta.tag}")
    document.add_paragraph(f"Commit: {meta.sha}")
    document.add_paragraph(f"Generated: {meta.generated_at}")


def _add_result_summary(document: docx.Document, doc: QaDocument) -> None:
    """Add the headline verdict, metrics, and configured quality gates list."""
    document.add_heading("Result Summary", level=1)
    reg = doc.regression
    document.add_paragraph(f"Verdict: {reg.verdict}")
    if reg.available:
        cov = f"{reg.coverage_pct}%" if reg.coverage_pct is not None else "N/A"
        document.add_paragraph(
            f"Tests: {reg.total} total, {reg.passed} passed, "
            f"{reg.failed} failed, {reg.skipped} skipped — Coverage: {cov}"
        )
    document.add_heading("Configured quality gates", level=2)
    for gate in doc.gates:
        document.add_paragraph(gate, style="List Bullet")


def _add_one_catalog_section(document: docx.Document, section: TestSection) -> None:
    """Render one TestSection as a heading, optional intro, and a table."""
    document.add_heading(f"{section.datetime} — {section.title}", level=2)
    if section.intro:
        document.add_paragraph(section.intro)
    table = document.add_table(rows=1, cols=len(CATALOG_COLUMNS))
    table.style = GRID_STYLE
    _fill_table_row(table.rows[0], CATALOG_COLUMNS)
    for case in section.cases:
        row = table.add_row()
        _fill_table_row(
            row,
            (
                str(case.number),
                case.description,
                case.type,
                case.test_file,
                case.expected,
            ),
        )


def _add_testcase_catalog(
    document: docx.Document, sections: tuple[TestSection, ...]
) -> None:
    """Add the full test-case catalog section."""
    document.add_heading("Test-case Catalog", level=1)
    for section in sections:
        _add_one_catalog_section(document, section)


def _add_regression_totals_table(
    document: docx.Document, regression: RegressionResult
) -> None:
    """Add the regression totals table when data is available."""
    headers = ("Tests", "Passed", "Failed", "Skipped", "Errors", "Duration (s)")
    table = document.add_table(rows=2, cols=len(headers))
    table.style = GRID_STYLE
    _fill_table_row(table.rows[0], headers)
    _fill_table_row(
        table.rows[1],
        (
            str(regression.total),
            str(regression.passed),
            str(regression.failed),
            str(regression.skipped),
            str(regression.errors),
            str(regression.duration_s),
        ),
    )
    cov = (
        f"{regression.coverage_pct}%"
        if regression.coverage_pct is not None
        else "N/A"
    )
    document.add_paragraph(f"Coverage: {cov}")


def _add_regression_failures(
    document: docx.Document, regression: RegressionResult
) -> None:
    """List individual test failures when present."""
    if not regression.failures:
        return
    document.add_heading("Failures", level=2)
    for fd in regression.failures:
        document.add_paragraph(f"{fd.test_id}: {fd.message}", style="List Bullet")


def _add_regression_results(
    document: docx.Document, regression: RegressionResult
) -> None:
    """Add the regression results section."""
    document.add_heading("Regression Results", level=1)
    if not regression.available:
        document.add_paragraph("Regression data unavailable.")
        return
    _add_regression_totals_table(document, regression)
    _add_regression_failures(document, regression)


def _add_signoff_block(document: docx.Document) -> None:
    """Add the sign-off signature table and release decision line."""
    document.add_heading("Sign-off", level=1)
    table = document.add_table(
        rows=1 + len(SIGNOFF_ROLES), cols=len(SIGNOFF_COLUMNS)
    )
    table.style = GRID_STYLE
    _fill_table_row(table.rows[0], SIGNOFF_COLUMNS)
    for row, role in zip(table.rows[1:], SIGNOFF_ROLES, strict=True):
        row.cells[0].text = role
    document.add_paragraph("Release decision:  ☐ Approved   ☐ Rejected")


def build_document(doc: QaDocument, out_dir: Path) -> Path:
    """Build a .docx QA sign-off report and return the saved file path."""
    document = docx.Document()
    _add_title_header(document, doc.meta)
    _add_result_summary(document, doc)
    _add_testcase_catalog(document, doc.sections)
    _add_regression_results(document, doc.regression)
    _add_signoff_block(document)
    filename = _output_filename(doc.meta)
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / filename
    document.save(str(out_path))
    return out_path
