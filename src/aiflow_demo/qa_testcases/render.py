"""Render a TestPlan into a Word (.docx) QA test-case document."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

import docx

from .models import ModuleCoverage, TestCaseRow, TestPlan

COVERAGE_COLUMNS = (
    "Module", "User Stories", "Total TCs",
    "Positive", "Negative", "Boundary", "Permission",
)
CASE_COLUMNS = (
    "User Story", "Acceptance Criteria", "Test Case (Module/Function)",
    "Test Steps", "Test Data", "Scenario Type", "Test Result (Pass/Fail & Reason)",
)
SIGNOFF_FIELDS = (
    "QA Engineer Name", "Date of Review", "Test Environment / Build Version",
    "Total Test Cases Executed", "Passed", "Failed", "Blocked / Not Applicable",
)
LEGEND = (
    "Positive — Happy path, expected behavior under normal conditions",
    "Negative — Invalid inputs, error handling, unexpected actions",
    "Boundary — Edge values (empty, zero, max, min, limits)",
    "Permission — Access control, role restrictions, immutability",
)
GRID_STYLE = "Table Grid"


@dataclass(frozen=True)
class RenderMeta:
    """Document metadata not carried by the plan itself."""

    branch: str
    generated_at: str
    date_stamp: str


def _fill_row(row: docx.table.Row, values: tuple[str, ...]) -> None:
    for cell, value in zip(row.cells, values, strict=False):
        cell.text = value


def _add_header(document: docx.Document, plan: TestPlan, meta: RenderMeta) -> None:
    document.add_heading(f"Test Case Document — {plan.project}", level=0)
    document.add_paragraph(f"Project: {plan.project} ({meta.branch})")
    document.add_paragraph(f"Generated: {meta.generated_at}")
    document.add_paragraph(f"Branch: {meta.branch}")
    document.add_paragraph("Prepared for: QA Team — Manual Testing")


def _add_summary(document: docx.Document, plan: TestPlan) -> None:
    document.add_heading("Summary", level=1)
    document.add_paragraph(plan.summary)
    document.add_heading("Assumptions", level=2)
    for item in plan.assumptions:
        document.add_paragraph(item, style="List Bullet")


def _coverage_values(cov: ModuleCoverage) -> tuple[str, ...]:
    return (
        cov.module, str(cov.user_stories), str(cov.total),
        str(cov.positive), str(cov.negative), str(cov.boundary), str(cov.permission),
    )


def _coverage_total(coverage: tuple[ModuleCoverage, ...]) -> tuple[str, ...]:
    return (
        "TOTAL",
        str(sum(c.user_stories for c in coverage)),
        str(sum(c.total for c in coverage)),
        str(sum(c.positive for c in coverage)),
        str(sum(c.negative for c in coverage)),
        str(sum(c.boundary for c in coverage)),
        str(sum(c.permission for c in coverage)),
    )


def _add_coverage_table(document: docx.Document, plan: TestPlan) -> None:
    document.add_heading("Test Cases", level=1)
    table = document.add_table(rows=1, cols=len(COVERAGE_COLUMNS))
    table.style = GRID_STYLE
    _fill_row(table.rows[0], COVERAGE_COLUMNS)
    for cov in plan.coverage:
        _fill_row(table.add_row(), _coverage_values(cov))
    _fill_row(table.add_row(), _coverage_total(plan.coverage))


def _add_legend(document: docx.Document) -> None:
    document.add_heading("Scenario Type Legend", level=2)
    for line in LEGEND:
        document.add_paragraph(line)


def _case_values(case: TestCaseRow) -> tuple[str, ...]:
    return (
        case.user_story, case.acceptance_criteria, case.test_case,
        case.test_steps, case.test_data, case.scenario_type, "",
    )


def _add_case_table(document: docx.Document, plan: TestPlan) -> None:
    table = document.add_table(rows=1, cols=len(CASE_COLUMNS))
    table.style = GRID_STYLE
    _fill_row(table.rows[0], CASE_COLUMNS)
    for case in plan.test_cases:
        _fill_row(table.add_row(), _case_values(case))


def _add_signoff(document: docx.Document) -> None:
    document.add_heading("QA Sign-Off", level=1)
    document.add_paragraph(
        "This section documents the QA review and confirmation of the test cases "
        "above. The QA engineer confirms that all applicable test cases have been "
        "executed and the results recorded."
    )
    table = document.add_table(rows=len(SIGNOFF_FIELDS), cols=2)
    table.style = GRID_STYLE
    for row, label in zip(table.rows, SIGNOFF_FIELDS, strict=True):
        row.cells[0].text = label
    document.add_heading("Remarks / Notes", level=2)
    document.add_paragraph(
        "(Any observations, deviations, or recommendations from the QA review)"
    )
    document.add_paragraph("QA Engineer Signature: ___________________________")
    document.add_paragraph("QA Lead / Manager Approval: ___________________________")
    document.add_paragraph("Date: ___________________________")


def _output_name(plan: TestPlan, meta: RenderMeta) -> str:
    safe = "".join(c if c.isalnum() else "_" for c in plan.project) or "Project"
    return f"Test_Cases_{safe}_{meta.date_stamp}.docx"


def render_document(plan: TestPlan, meta: RenderMeta, out_dir: Path) -> Path:
    """Render the plan to a .docx under out_dir and return the saved path."""
    document = docx.Document()
    _add_header(document, plan, meta)
    _add_summary(document, plan)
    _add_coverage_table(document, plan)
    _add_legend(document)
    _add_case_table(document, plan)
    _add_signoff(document)
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / _output_name(plan, meta)
    document.save(str(out_path))
    return out_path
